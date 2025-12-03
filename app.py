import streamlit as st
import pandas as pd
from PyPDF2 import PdfReader
from docx import Document
import google.generativeai as genai
import io
import time
import zipfile
import json
import re
import sqlite3
import hashlib
from datetime import datetime
from fpdf import FPDF
import plotly.express as px

# ==============================================================================
# 1. CONFIGURACIÓN E INICIALIZACIÓN
# ==============================================================================

st.set_page_config(
    page_title="HR Intelligence Suite",
    layout="wide",
    page_icon="🎓",
    initial_sidebar_state="expanded"
)

# Inicializar Base de Datos SQLite
def init_db():
    conn = sqlite3.connect('cv_master_db.db', check_same_thread=False)
    c = conn.cursor()
    # Tabla única robusta
    c.execute('''CREATE TABLE IF NOT EXISTS analisis (
                    file_hash TEXT PRIMARY KEY,
                    timestamp TEXT,
                    lote_nombre TEXT,
                    filename TEXT,
                    candidato TEXT,
                    facultad TEXT,
                    cargo TEXT,
                    puntaje REAL,
                    recomendacion TEXT,
                    ajuste TEXT,
                    raw_json TEXT,
                    pdf_blob BLOB
                )''')
    conn.commit()
    return conn

conn = init_db()

# --- Funciones de Utilidad BD ---
def get_file_hash(file_bytes):
    return hashlib.md5(file_bytes).hexdigest()

def db_exists(file_hash):
    c = conn.cursor()
    c.execute("SELECT candidato FROM analisis WHERE file_hash = ?", (file_hash,))
    return c.fetchone() is not None

def db_save(data, pdf_bytes, file_hash, filename, lote):
    c = conn.cursor()
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    json_str = json.dumps(data, ensure_ascii=False)
    
    # Upsert (Insertar o Reemplazar si ya existe el hash)
    c.execute('''INSERT OR REPLACE INTO analisis 
                 VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
              (file_hash, now, lote, filename, 
               data.get('nombre', 'N/A'),
               data.get('facultad', ''),
               data.get('cargo', ''),
               data.get('puntaje_global', 0.0),
               data.get('recomendacion', 'N/A'),
               data.get('ajuste', 'N/A'),
               json_str, pdf_bytes))
    conn.commit()

def db_load_all():
    return pd.read_sql("SELECT * FROM analisis ORDER BY timestamp DESC", conn)

# ==============================================================================
# 2. MOTORES DE LECTURA Y AI (ROBUSTEZ)
# ==============================================================================

def read_file_safe(file_obj):
    """Lee el archivo asegurando el puntero al inicio."""
    try:
        file_obj.seek(0)
        if file_obj.type == "application/pdf":
            reader = PdfReader(file_obj)
            return "".join([page.extract_text() or "" for page in reader.pages])
        elif "word" in file_obj.type or file_obj.name.endswith(".docx"):
            doc = Document(file_obj)
            return "\n".join([p.text for p in doc.paragraphs])
        return ""
    except Exception: return ""

def get_dynamic_model(api_key):
    """
    AUTODESCUBRIMIENTO: Consulta a la API qué modelos están disponibles
    para evitar errores de 'Model not found'.
    """
    genai.configure(api_key=api_key)
    try:
        # Listar modelos que soporten generación de texto
        available = []
        for m in genai.list_models():
            if 'generateContent' in m.supported_generation_methods:
                available.append(m.name)
        
        # Lógica de priorización inteligente
        # 1. Buscar Flash 1.5 o 2.0 (Rapidez y Costo)
        for m in available:
            if 'flash' in m.lower() and ('1.5' in m or '2.0' in m): return m
        # 2. Buscar Pro 1.5 (Potencia)
        for m in available:
            if 'pro' in m.lower() and '1.5' in m: return m
        # 3. Fallback a cualquiera disponible
        if available: return available[0]
        
        return 'models/gemini-1.5-flash' # Default ciego si falla el listado
    except:
        return 'models/gemini-1.5-flash'

def analyze_cv(text, role, faculty, api_key):
    if not api_key: return None
    
    # 1. Obtener modelo válido
    model_name = get_dynamic_model(api_key)
    genai.configure(api_key=api_key)
    
    # 2. Rúbrica
    base_rubrica = {"Formación": "35%", "Experiencia": "30%", "Competencias": "20%", "Software": "15%"}
    
    # 3. Prompt Estricto
    prompt = f"""
    Actúa como Experto en Selección Académica. Evalúa CV para Cargo: "{role}" en Facultad: "{faculty}".
    RÚBRICA: {json.dumps(base_rubrica)}
    
    INSTRUCCIONES CRÍTICAS:
    1. Responde SOLO JSON válido.
    2. "puntaje_global": float con 2 decimales (1.00 a 5.00).
    3. "recomendacion": "NO RECOMENDADO", "REQUIERE ANTECEDENTES", o "AVANZA".
    
    JSON:
    {{
        "nombre": "Nombre Apellido",
        "ajuste": "ALTO/MEDIO/BAJO",
        "puntaje_global": 0.00,
        "recomendacion": "ESTADO",
        "conclusion_ejecutiva": "Resumen.",
        "detalle": {{
            "formacion": {{ "nota": 0, "ponderado": 0.00 }},
            "experiencia": {{ "nota": 0, "ponderado": 0.00 }},
            "competencias": {{ "nota": 0, "ponderado": 0.00 }},
            "software": {{ "nota": 0, "ponderado": 0.00 }}
        }},
        "cualitativo": {{
            "brechas": ["..."],
            "riesgos": ["..."],
            "fortalezas": ["..."]
        }}
    }}
    CV: {text[:30000]}
    """
    
    try:
        model = genai.GenerativeModel(model_name)
        response = model.generate_content(prompt)
        
        # Limpieza JSON (Find first { and last })
        raw = response.text
        start = raw.find('{')
        end = raw.rfind('}') + 1
        if start == -1: return None
        
        return json.loads(raw[start:end])
    except Exception as e:
        print(f"Error AI ({model_name}): {e}")
        return None

# ==============================================================================
# 3. GENERADOR PDF (DISEÑO SOLICITADO)
# ==============================================================================

class PDFReport(FPDF):
    def header(self):
        self.set_font('Arial', 'B', 14)
        self.cell(0, 10, 'INFORME DE AJUSTE CANDIDATO-CARGO', 0, 1, 'C')
        self.ln(10)
    def footer(self):
        self.set_y(-15); self.set_font('Arial', 'I', 8)
        self.cell(0, 10, f'HR Suite - {datetime.now().strftime("%Y-%m-%d")}', 0, 0, 'C')

def generate_pdf(data):
    pdf = PDFReport()
    pdf.add_page()
    def txt(s): return str(s).encode('latin-1', 'replace').decode('latin-1')
    
    # Datos
    pdf.set_font('Arial', 'B', 11)
    pdf.cell(30, 6, "Candidato:", 0, 0); pdf.set_font('Arial', '', 11)
    pdf.cell(0, 6, txt(data.get('nombre', '')), 0, 1)
    pdf.set_font('Arial', 'B', 11)
    pdf.cell(30, 6, "Cargo:", 0, 0); pdf.set_font('Arial', '', 11)
    pdf.cell(0, 6, txt(f"{data.get('cargo')} / {data.get('facultad')}"), 0, 1)
    pdf.ln(5)
    
    # A. Conclusión
    pdf.set_font('Arial', 'B', 12); pdf.set_fill_color(240, 240, 240)
    pdf.cell(0, 8, txt("A. CONCLUSIÓN EJECUTIVA"), 1, 1, 'L', True); pdf.ln(2)
    pdf.set_font('Arial', '', 10)
    pdf.multi_cell(0, 5, txt(f"Ajuste: {data.get('ajuste')}. Puntaje: {data.get('puntaje_global')}/5.00.\n{data.get('conclusion_ejecutiva')}"))
    pdf.ln(5)
    
    # Estado (Color)
    rec = data.get('recomendacion', '').upper()
    if "NO" in rec: pdf.set_fill_color(255, 200, 200)
    elif "AVANZA" in rec: pdf.set_fill_color(200, 255, 200)
    else: pdf.set_fill_color(255, 255, 200)
    pdf.set_font('Arial', 'B', 12)
    pdf.cell(0, 10, txt(rec), 1, 1, 'C', True); pdf.ln(8)
    
    # B. Tabla
    pdf.set_font('Arial', 'B', 12); pdf.set_fill_color(240, 240, 240)
    pdf.cell(0, 8, txt("B. TABLA RESUMEN"), 1, 1, 'L', True)
    pdf.set_font('Arial', 'B', 9); pdf.set_fill_color(50, 50, 50); pdf.set_text_color(255)
    pdf.cell(80, 8, "Variable", 1, 0, 'L', True); pdf.cell(30, 8, "Ponderacion", 1, 0, 'C', True)
    pdf.cell(30, 8, "Nota", 1, 0, 'C', True); pdf.cell(50, 8, "Puntaje", 1, 1, 'C', True)
    pdf.set_text_color(0); pdf.set_font('Arial', '', 9)
    
    det = data.get('detalle', {})
    dims = [("Formación", "35%", det.get('formacion', {})), ("Experiencia", "30%", det.get('experiencia', {})),
            ("Competencias", "20%", det.get('competencias', {})), ("Software", "15%", det.get('software', {}))]
    
    for n, p, v in dims:
        pdf.ln(8)
        pdf.cell(80, 8, txt(n), 1)
        pdf.cell(30, 8, p, 1, 0, 'C')
        pdf.cell(30, 8, str(v.get('nota', 0)), 1, 0, 'C')
        pdf.cell(50, 8, f"{v.get('ponderado', 0):.2f}", 1, 0, 'C')
    pdf.ln(10)
    
    # C. Cualitativo
    pdf.set_font('Arial', 'B', 12); pdf.set_fill_color(240, 240, 240)
    pdf.cell(0, 8, txt("C. ANÁLISIS CUALITATIVO"), 1, 1, 'L', True); pdf.ln(2)
    pdf.set_font('Arial', '', 10)
    cual = data.get('cualitativo', {})
    for k, v in cual.items():
        pdf.set_font('Arial', 'B', 10); pdf.cell(0, 6, txt(k.upper()), 0, 1)
        pdf.set_font('Arial', '', 10)
        items = v if isinstance(v, list) else [str(v)]
        for i in items: pdf.multi_cell(0, 5, txt(f"- {i}"))
        pdf.ln(2)
        
    return bytes(pdf.output())

# ==============================================================================
# 4. LÓGICA DE PROCESAMIENTO (LINEAR)
# ==============================================================================

def process_batch_files(file_list, faculty, role, lote_name, api_key, skip_dupes, progress_bar, status_text):
    """Procesa una lista de archivos secuencialmente actualizando la UI."""
    count_ok = 0
    total = len(file_list)
    
    for i, file in enumerate(file_list):
        status_text.text(f"Procesando {i+1}/{total}: {file.name}...")
        
        # 1. Leer Binario para Hash
        file.seek(0)
        file_bytes = file.read()
        file_hash = get_file_hash(file_bytes)
        
        # 2. Verificar Duplicidad
        if skip_dupes and db_exists(file_hash):
            continue # Saltar silenciosamente
            
        # 3. Leer Texto
        file.seek(0) # Reset puntero para lector PDF/Docx
        text = read_file_safe(file)
        
        if len(text) > 50:
            # 4. Analizar
            res = analyze_cv(text, role, faculty, api_key)
            if res:
                # 5. Enriquecer datos
                res['facultad'] = faculty
                res['cargo'] = role
                res['nombre'] = res.get('nombre', 'Candidato')
                
                # 6. Generar PDF y Guardar
                pdf_bytes = generate_pdf(res)
                db_save(res, pdf_bytes, file_hash, file.name, lote_name)
                count_ok += 1
        
        # Actualizar barra
        progress_bar.progress((i + 1) / total)
        
    return count_ok

# ==============================================================================
# 5. INTERFAZ DE USUARIO (CLEAN & NATIVE)
# ==============================================================================

# --- Sidebar ---
with st.sidebar:
    st.title("⚙️ Panel de Control")
    
    # API Key
    api_key = st.text_input("Google API Key", type="password", help="Obtener en aistudio.google.com")
    if 'GOOGLE_API_KEY' in st.secrets:
        api_key = st.secrets['GOOGLE_API_KEY']
        st.success("Llave Institucional Detectada")
        
    st.divider()
    skip_dupes = st.checkbox("Omitir Duplicados", value=True, help="No vuelve a analizar archivos idénticos.")
    
    if st.button("🔴 Borrar Todo el Historial"):
        conn.cursor().execute("DELETE FROM analisis")
        conn.commit()
        st.warning("Base de datos vaciada.")
        time.sleep(1)
        st.rerun()

# --- Main ---
st.title("🚀 HR Intelligence Suite")
st.markdown("Sistema integral de análisis curricular.")

# Pestañas
tab_load, tab_dash, tab_data = st.tabs(["📥 Centro de Carga", "📊 Dashboard", "🗃️ Base de Datos"])

# --- TAB 1: CARGA ---
with tab_load:
    # Contenedores Nativos (Mejor que tarjetas HTML personalizadas que rompen contraste)
    col1, col2 = st.columns(2)
    
    batches_to_run = []
    
    # Definición de Facultades Correcta
    FACULTADES = [
        "Facultad de Ingeniería",
        "Facultad de Economía y Negocios",
        "Facultad de Ciencias de la Vida",
        "Facultad de Educación y Ciencias Sociales"
    ]
    CARGOS = ["Docente", "Investigador", "Gestión Académica"]

    def render_batch_ui(column, idx):
        with column:
            with st.container(border=True):
                st.subheader(f"📂 Lote #{idx}")
                files = st.file_uploader(f"CVs Lote {idx}", type=['pdf','docx'], accept_multiple_files=True, key=f"u{idx}")
                fac = st.selectbox("Facultad", FACULTADES, key=f"f{idx}")
                rol = st.selectbox("Cargo", CARGOS, key=f"r{idx}")
                
                # Botón Individual
                if st.button(f"▶ Procesar Lote {idx}", key=f"b{idx}"):
                    if not api_key: st.error("Falta API Key")
                    elif not files: st.warning("Sin archivos")
                    else:
                        bar = st.progress(0)
                        stat = st.empty()
                        n = process_batch_files(files, fac, rol, f"Lote {idx}", api_key, skip_dupes, bar, stat)
                        st.success(f"Procesados: {n}")
                        time.sleep(1)
                        st.rerun()
                
                if files:
                    batches_to_run.append({'files': files, 'fac': fac, 'rol': rol, 'id': f"Lote {idx}"})

    render_batch_ui(col1, 1)
    render_batch_ui(col2, 2)
    render_batch_ui(col1, 3)
    render_batch_ui(col2, 4)
    
    st.markdown("---")
    
    # Botón Global
    total_files_global = sum([len(b['files']) for b in batches_to_run])
    
    if st.button("🚀 PROCESAR TODOS LOS LOTES ACTIVOS", type="primary", use_container_width=True):
        if not api_key:
            st.error("Falta API Key")
        elif total_files_global == 0:
            st.warning("No hay archivos seleccionados en ningún lote.")
        else:
            # Estimación
            mins = (total_files_global * 5) // 60
            st.info(f"Iniciando proceso para {total_files_global} documentos. Tiempo estimado: {mins} min aprox.")
            
            global_bar = st.progress(0)
            global_stat = st.empty()
            
            total_processed = 0
            current_idx = 0
            
            for batch in batches_to_run:
                # Proceso manual del batch para actualizar la barra global
                for f in batch['files']:
                    global_stat.text(f"Analizando: {f.name} ({batch['id']})")
                    
                    # Lógica de Hash y Proceso
                    f.seek(0); f_bytes = f.read(); f_hash = get_file_hash(f_bytes)
                    
                    if not (skip_dupes and db_exists(f_hash)):
                        f.seek(0)
                        text = read_file_safe(f)
                        if len(text) > 50:
                            res = analyze_cv(text, batch['rol'], batch['fac'], api_key)
                            if res:
                                res.update({'nombre': res.get('nombre', 'N/A'), 'facultad': batch['fac'], 'cargo': batch['rol']})
                                pdf = generate_pdf(res)
                                db_save(res, pdf, f_hash, f.name, batch['id'])
                                total_processed += 1
                    
                    current_idx += 1
                    global_bar.progress(current_idx / total_files_global)
            
            st.success(f"¡Listo! Se procesaron {total_processed} nuevos documentos.")
            time.sleep(2)
            st.rerun()

# --- TAB 2: DASHBOARD ---
with tab_dash:
    st.subheader("Tablero de Control")
    df = db_load_all()
    
    if not df.empty:
        # KPIs
        k1, k2, k3 = st.columns(3)
        k1.metric("Total Evaluados", len(df))
        avg = df['puntaje'].mean()
        k2.metric("Promedio Global", f"{avg:.2f}")
        aprob = len(df[df['recomendacion'].str.contains("AVANZA", case=False, na=False)])
        k3.metric("Candidatos Aptos", aprob)
        
        st.divider()
        
        # Gráficos
        c_chart1, c_chart2 = st.columns(2)
        
        with c_chart1:
            st.markdown("##### Distribución por Facultad")
            if 'facultad' in df.columns and 'puntaje' in df.columns:
                fig = px.box(df, x="facultad", y="puntaje", color="facultad")
                st.plotly_chart(fig, use_container_width=True)
                
        with c_chart2:
            st.markdown("##### Estado de Candidatos")
            if 'recomendacion' in df.columns:
                fig2 = px.pie(df, names="recomendacion", hole=0.4)
                st.plotly_chart(fig2, use_container_width=True)
                
        # Ranking
        st.markdown("##### Top 10 Candidatos")
        st.dataframe(
            df.sort_values(by='puntaje', ascending=False).head(10)[['candidato', 'puntaje', 'recomendacion', 'facultad']],
            use_container_width=True,
            hide_index=True
        )
    else:
        st.info("Aún no hay datos para mostrar.")

# --- TAB 3: BASE DE DATOS ---
with tab_data:
    st.subheader("Registro Histórico Completo")
    df = db_load_all()
    
    if not df.empty:
        # Descarga Masiva ZIP
        zip_mem = io.BytesIO()
        with zipfile.ZipFile(zip_mem, "w") as zf:
            for i, row in df.iterrows():
                if row['pdf_blob']:
                    name = re.sub(r'[^a-zA-Z0-9]', '_', str(row['candidato']))
                    zf.writestr(f"{name}.pdf", row['pdf_blob'])
        
        c_d1, c_d2 = st.columns([1, 2])
        c_d1.download_button("📦 Descargar Todos los PDFs (ZIP)", zip_mem.getvalue(), "Informes.zip", "application/zip", type="primary")
        c_d2.download_button("💾 Descargar Excel", df.drop(columns=['pdf_blob', 'raw_json']).to_csv(index=False).encode('utf-8'), "data.csv", "text/csv")
        
        st.dataframe(df.drop(columns=['pdf_blob', 'raw_json']), use_container_width=True)
    else:
        st.info("Base de datos vacía.")
